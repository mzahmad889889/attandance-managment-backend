from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required
from src.models.attendance_model import AttendanceRecord
from src.models.worker_model import Worker
from src.models.plant_model import Plant
from src.models.contractor_model import Contractor
from datetime import date, datetime, timedelta
import io, os

report_bp = Blueprint('reports', __name__)


@report_bp.route('/export-excel', methods=['GET'])
@jwt_required()
def export_excel():
    """Export attendance to Excel using pandas + openpyxl."""
    try:
        import pandas as pd
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return jsonify({'error': 'pandas/openpyxl not installed'}), 503

    date_from_str = request.args.get('date_from')
    date_to_str = request.args.get('date_to')
    plant_id = request.args.get('plant_id', type=int)
    contractor_id = request.args.get('contractor_id', type=int)

    # Default: current month
    today = date.today()
    if date_from_str:
        try:
            date_from = date.fromisoformat(date_from_str)
        except ValueError:
            date_from = date(today.year, today.month, 1)
    else:
        date_from = date(today.year, today.month, 1)

    if date_to_str:
        try:
            date_to = date.fromisoformat(date_to_str)
        except ValueError:
            date_to = today
    else:
        date_to = today

    q = AttendanceRecord.query.join(Worker).filter(
        AttendanceRecord.date >= date_from,
        AttendanceRecord.date <= date_to,
        Worker.is_active == True,
    )
    if plant_id:
        q = q.filter(Worker.plant_id == plant_id)
    if contractor_id:
        q = q.filter(Worker.contractor_id == contractor_id)

    records = q.order_by(AttendanceRecord.date, Worker.name).all()

    rows = []
    for r in records:
        w = r.worker
        rows.append({
            'Date': r.date.isoformat() if r.date else '',
            'Worker Code': w.worker_code if w else '',
            'Name': w.name if w else '',
            'Plant': w.plant.name if w and w.plant else '',
            'Contractor': w.contractor.name if w and w.contractor else '',
            'Shift': r.shift_type or '',
            'Check In': r.checkin_time.strftime('%H:%M') if r.checkin_time else '',
            'Check Out': r.checkout_time.strftime('%H:%M') if r.checkout_time else '',
            'Total Hours': r.total_hours or 0,
            'Overtime Hours': r.overtime_hours or 0,
            'Status': r.status or '',
            'Live Status': r.live_status or '',
        })

    df = pd.DataFrame(rows)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Attendance', index=False)
        ws = writer.sheets['Attendance']
        # Basic styling
        header_fill = PatternFill(start_color='F97316', end_color='F97316', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for col_idx, col in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        for col_idx in range(1, len(df.columns) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 16

    buf.seek(0)
    filename = f'attendance_{date_from}_{date_to}.xlsx'
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


@report_bp.route('/summary', methods=['GET'])
@jwt_required()
def summary():
    """Weekly and monthly summary stats."""
    today = date.today()
    week_start = today - timedelta(days=today.weekday())

    # Last 7 days chart data
    chart_data = []
    for i in range(6, -1, -1):
        d = today - timedelta(days=i)
        records = AttendanceRecord.query.filter_by(date=d).all()
        present = sum(1 for r in records if r.status in ('Present', 'Late'))
        total = Worker.query.filter_by(is_active=True).count()
        chart_data.append({
            'date': d.isoformat(),
            'day': d.strftime('%a'),
            'present': present,
            'absent': total - present,
        })

    # Plant breakdown
    plants = Plant.query.all()
    plant_data = []
    for p in plants:
        worker_ids = [w.id for w in p.workers if w.is_active]
        today_in = AttendanceRecord.query.filter(
            AttendanceRecord.worker_id.in_(worker_ids),
            AttendanceRecord.date == today,
            AttendanceRecord.live_status == 'IN'
        ).count() if worker_ids else 0
        plant_data.append({
            'plant': p.name,
            'total': len(worker_ids),
            'active_now': today_in,
            'capacity': p.capacity,
        })

    # Monthly overtime
    month_start = date(today.year, today.month, 1)
    monthly_ot = AttendanceRecord.query.filter(
        AttendanceRecord.date >= month_start,
        AttendanceRecord.date <= today
    ).with_entities(
        db.func.sum(AttendanceRecord.overtime_hours)
    ).scalar() or 0

    return jsonify({
        'chart_data': chart_data,
        'plant_breakdown': plant_data,
        'monthly_overtime_hours': round(float(monthly_ot), 2),
    }), 200


@report_bp.route('/worker/<int:worker_id>/history', methods=['GET'])
@jwt_required()
def worker_history(worker_id):
    """Fetch 1 month history for a specific worker."""
    limit = date.today() - timedelta(days=30)
    records = AttendanceRecord.query.filter(
        AttendanceRecord.worker_id == worker_id,
        AttendanceRecord.date >= limit
    ).order_by(AttendanceRecord.date.desc()).all()
    
    return jsonify({
        'history': [r.to_dict() for r in records]
    }), 200


@report_bp.route('/worker/<int:worker_id>/export', methods=['GET'])
@jwt_required()
def export_worker_excel(worker_id):
    """Export 1 month history for a single worker to Excel."""
    try:
        import pandas as pd
    except ImportError:
        return jsonify({'error': 'pandas not installed'}), 503
        
    worker = Worker.query.get_or_404(worker_id)
    limit = date.today() - timedelta(days=30)
    records = AttendanceRecord.query.filter(
        AttendanceRecord.worker_id == worker_id,
        AttendanceRecord.date >= limit
    ).order_by(AttendanceRecord.date.asc()).all()
    
    rows = []
    for r in records:
        rows.append({
            'Date': r.date.isoformat(),
            'Shift': r.shift_type,
            'In': r.checkin_time.strftime('%H:%M') if r.checkin_time else '',
            'Out': r.checkout_time.strftime('%H:%M') if r.checkout_time else '',
            'Total Hrs': r.total_hours,
            'OT Hrs': r.overtime_hours,
            'Status': r.status
        })
    
    df = pd.DataFrame(rows)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name=f'History_{worker.worker_code}', index=False)
    
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'History_{worker.worker_code}_{date.today()}.xlsx'
    )


# Local import needed for db
from src.extention import db
